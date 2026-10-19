"""Um Drive falso para os testes da T6 (sem rede, sem Google). Duas peças:

- `DriveFalso`: no lugar do LeitorDrive (a importação, o script e a API só usam `arquivos`, `baixar` e `requisicoes`).
  Guarda arquivos de verdade (bytes de docx, PDF, xlsx...) e anota o que foi baixado.
- `SessaoFalsa`: no lugar da sessão HTTP do google-auth, para testar o LeitorDrive em si (listar, paginar, baixar).

E construtores de arquivos reais em memória (`docx_com`, `pdf_com`, `xlsx_com`) para a extração de texto rodar nos
formatos de verdade. Os textos dos testes são inventados: nenhum documento do CITi entra no repositório.
"""
import io
import zlib
from datetime import datetime
from pathlib import Path

from app.drive import API, MIME_ATALHO, MIME_PASTA, ArquivoDrive, BuscaDrive, ErroDrive
from app.extrair_texto import MIME_DOCX, MIME_PDF, MIME_XLSX

BASE = datetime(2026, 10, 13, 12, 0)


# --- arquivos de verdade, em memória ------------------------------------------------------------------------

def docx_com(*blocos: str) -> bytes:
    """Um .docx. Bloco "# Título" vira título; "- item" vira lista; "| a | b |" começa uma tabela de uma linha."""
    from docx import Document

    doc = Document()
    for bloco in blocos:
        if bloco.startswith("# "):
            doc.add_heading(bloco[2:], level=1)
        elif bloco.startswith("## "):
            doc.add_heading(bloco[3:], level=2)
        elif bloco.startswith("- "):
            doc.add_paragraph(bloco[2:], style="List Bullet")
        elif bloco.startswith("| "):
            celulas = [c.strip() for c in bloco.strip("| ").split("|")]
            tabela = doc.add_table(rows=1, cols=len(celulas))
            for i, c in enumerate(celulas):
                tabela.rows[0].cells[i].text = c
        else:
            doc.add_paragraph(bloco)
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def xlsx_com(linhas: list[list], aba: str = "Página1") -> bytes:
    from openpyxl import Workbook

    pasta = Workbook()
    pasta.active.title = aba
    for linha in linhas:
        pasta.active.append(linha)
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def pdf_com(*linhas: str) -> bytes:
    """Um PDF de uma página com texto (fonte Helvetica, codificação WinAnsi: acentos funcionam). Sem linhas, a
    página fica em branco: é o PDF "de imagem", que não tem camada de texto."""
    def escapar(texto: str) -> bytes:
        return texto.encode("cp1252").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")

    conteudo = b"".join(
        b"BT /F1 11 Tf 56 %d Td (%s) Tj ET\n" % (780 - 16 * i, escapar(linha)) for i, linha in enumerate(linhas)
    )
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(zlib.compress(conteudo))
        + zlib.compress(conteudo) + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    saida = bytearray(b"%PDF-1.4\n")
    posicoes = []
    for i, objeto in enumerate(objetos, start=1):
        posicoes.append(len(saida))
        saida += b"%d 0 obj\n%s\nendobj\n" % (i, objeto)
    xref = len(saida)
    saida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    saida += b"".join(b"%010d 00000 n \n" % p for p in posicoes)
    saida += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objetos) + 1, xref)
    return bytes(saida)


# --- o Drive falso (no lugar do LeitorDrive) -------------------------------------------------------------------

class DriveFalso:
    def __init__(self):
        self.conteudos: dict[str, bytes] = {}
        self.meta: dict[str, ArquivoDrive] = {}
        self.baixados: list[str] = []  # os ids baixados, na ordem (a prova de que o que não mudou não é baixado)
        self.falhar_download: set[str] = set()
        self.requisicoes = 0
        self.lista_incompleta = False
        self._n = 0

    def adicionar(self, nome: str, conteudo: bytes, mime: str | None = None, caminho: tuple[str, ...] = (),
                  modificado_em: datetime | None = None, tamanho: int | None = None) -> ArquivoDrive:
        self._n += 1
        mime = mime or _mime_do_nome(nome)
        arquivo = ArquivoDrive(
            id=f"arq{self._n:03d}", nome=nome, mime=mime, modificado_em=modificado_em or BASE,
            tamanho=len(conteudo) if tamanho is None else tamanho, caminho=caminho,
        )
        self.meta[arquivo.id] = arquivo
        self.conteudos[arquivo.id] = conteudo
        return arquivo

    def editar(self, arquivo: ArquivoDrive, conteudo: bytes, modificado_em: datetime) -> ArquivoDrive:
        novo = ArquivoDrive(**{**arquivo.__dict__, "modificado_em": modificado_em, "tamanho": len(conteudo)})
        self.meta[arquivo.id] = novo
        self.conteudos[arquivo.id] = conteudo
        return novo

    def tocar(self, arquivo: ArquivoDrive, modificado_em: datetime) -> None:
        """Muda só a data de edição (alguém abriu e salvou sem mudar nada)."""
        self.meta[arquivo.id] = ArquivoDrive(**{**arquivo.__dict__, "modificado_em": modificado_em})

    def remover(self, arquivo: ArquivoDrive) -> None:
        self.meta.pop(arquivo.id)
        self.conteudos.pop(arquivo.id)

    # --- o que a importação usa ---------------------------------------------------------------------------

    def arquivos(self, limite: int | None = None) -> BuscaDrive:
        self.requisicoes += 1
        lista = sorted(self.meta.values(), key=lambda a: (a.modificado_em, a.id))
        completa = not self.lista_incompleta and (limite is None or len(lista) <= limite)
        return BuscaDrive(arquivos=lista[:limite] if limite else lista, completa=completa)

    def baixar(self, arquivo: ArquivoDrive, destino: Path) -> None:
        self.requisicoes += 1
        self.baixados.append(arquivo.id)
        if arquivo.id in self.falhar_download:
            raise ErroDrive("O Drive devolveu um erro (500).", 500)
        destino.write_bytes(self.conteudos[arquivo.id])


def _mime_do_nome(nome: str) -> str:
    return {
        "docx": MIME_DOCX, "pdf": MIME_PDF, "xlsx": MIME_XLSX, "txt": "text/plain", "md": "text/markdown",
        "csv": "text/csv", "png": "image/png",
    }[nome.rsplit(".", 1)[-1]]


# --- a sessão HTTP falsa (no lugar da AuthorizedSession), para testar o LeitorDrive em si ----------------------

class RespostaFalsa:
    def __init__(self, status: int = 200, json: dict | None = None, corpo: bytes = b""):
        self.status_code = status
        self._json = json
        self._corpo = corpo
        self.fechada = False

    def json(self) -> dict:
        return self._json

    def iter_content(self, chunk_size: int):
        for i in range(0, len(self._corpo), chunk_size):
            yield self._corpo[i: i + chunk_size]

    def close(self) -> None:
        self.fechada = True


class SessaoFalsa:
    """Responde a GET /drive/v3/files (a lista, por pasta e por página) e a /files/{id} e /files/{id}/export."""

    def __init__(self):
        self.pastas: dict[str, list[list[dict]]] = {}  # id da pasta -> páginas de filhos
        self.corpos: dict[str, bytes] = {}  # id do arquivo (ou id/export) -> bytes
        self.pedidos: list[tuple[str, dict]] = []
        self.programadas: list = []  # respostas (ou exceções) usadas antes das normais, na ordem

    def get(self, url: str, params: dict | None = None, stream: bool = False, timeout: int | None = None):
        params = params or {}
        self.pedidos.append((url, params))
        if self.programadas:
            programada = self.programadas.pop(0)
            if isinstance(programada, Exception):
                raise programada
            return programada
        if url == API:
            pasta = params["q"].split("'")[1]
            paginas = self.pastas.get(pasta, [[]])
            indice = int(params.get("pageToken", 0))
            corpo = {"files": paginas[indice]}
            if indice + 1 < len(paginas):
                corpo["nextPageToken"] = str(indice + 1)
            return RespostaFalsa(json=corpo)
        chave = url.removeprefix(API + "/")
        return RespostaFalsa(corpo=self.corpos[chave])


def item(id_: str, nome: str, mime: str, modificado: str = "2026-10-13T07:59:22.000Z", tamanho: int | None = 100,
         md5: str | None = None) -> dict:
    """Um arquivo no formato da resposta do Drive."""
    dados = {"id": id_, "name": nome, "mimeType": mime, "modifiedTime": modificado}
    if tamanho is not None:
        dados["size"] = str(tamanho)
    if md5:
        dados["md5Checksum"] = md5
    return dados


def pasta(id_: str, nome: str) -> dict:
    return item(id_, nome, MIME_PASTA, tamanho=None)


def atalho(id_: str, nome: str) -> dict:
    return item(id_, nome, MIME_ATALHO, tamanho=None)
