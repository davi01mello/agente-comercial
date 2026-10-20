"""Arquivo do Drive → texto (markdown simples). Sem rede e sem modelo.

Formatos: docx, pdf (só a camada de texto; PDF de imagem dá texto vazio), xlsx, csv, txt e md. Os documentos nativos do
Google chegam aqui já exportados (texto puro ou CSV, pelo app/drive.py).

Devolve o texto; quem chama decide o que fazer com texto vazio ("sem texto": PDF de slides em imagem, planilha em
branco). Arquivo corrompido levanta ErroDeLeitura (só o tipo do erro: a mensagem de uma biblioteca pode ter texto do
arquivo).
"""
import csv
import re
from pathlib import Path

MAX_CARACTERES = 600_000  # um livro: acima disto o resto é cortado (e a busca ainda acha o que veio antes)
MAX_LINHAS_PLANILHA = 5_000
MAX_PAGINAS_PDF = 400  # acima disto o resto é ignorado (o texto de 400 páginas já passa do limite de caracteres)

MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_PDF = "application/pdf"
MIME_TEXTO = {"text/plain", "text/markdown", "text/x-markdown"}
MIME_CSV = {"text/csv", "application/csv"}
MIME_GOOGLE_PLANILHA = "application/vnd.google-apps.spreadsheet"
MIME_GOOGLE_TEXTO = {"application/vnd.google-apps.document", "application/vnd.google-apps.presentation"}

PLANILHAS = {MIME_XLSX, MIME_GOOGLE_PLANILHA, *MIME_CSV}  # não vão ao curador: lista de dados não é aprendizado
SUPORTADOS = {MIME_DOCX, MIME_XLSX, MIME_PDF, *MIME_TEXTO, *MIME_CSV, MIME_GOOGLE_PLANILHA, *MIME_GOOGLE_TEXTO}

_LINHAS_EM_BRANCO = re.compile(r"\n{3,}")
_ESPACOS_NA_PONTA = re.compile(r"[ \t]+\n")
_ESPACOS_REPETIDOS = re.compile(r"[ \t]{2,}")  # o docx exportado do Google vem com dezenas de espaços de enchimento


class ErroDeLeitura(Exception):
    """O arquivo está corrompido, protegido ou não é do tipo que diz ser. A mensagem é só o tipo do erro."""


def suportado(mime: str) -> bool:
    return mime in SUPORTADOS


def e_planilha(mime: str) -> bool:
    return mime in PLANILHAS


def extrair(caminho: Path, mime: str) -> str:
    """O texto do arquivo em `caminho`, de acordo com o `mime` do Drive."""
    try:
        if mime == MIME_DOCX:
            texto = _docx(caminho)
        elif mime == MIME_PDF:
            texto = _pdf(caminho)
        elif mime == MIME_XLSX:
            texto = _xlsx(caminho)
        elif mime in MIME_CSV or mime == MIME_GOOGLE_PLANILHA:
            texto = _csv(caminho)
        elif mime in MIME_TEXTO or mime in MIME_GOOGLE_TEXTO:
            texto = caminho.read_bytes().decode("utf-8-sig", errors="replace")
        else:
            raise ErroDeLeitura(f"formato não suportado ({mime})")
    except ErroDeLeitura:
        raise
    except Exception as e:  # as bibliotecas levantam de tudo (zip ruim, xml ruim, PDF cifrado...)
        raise ErroDeLeitura(type(e).__name__) from e
    return _limpar(texto)[:MAX_CARACTERES]


def _limpar(texto: str) -> str:
    texto = texto.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "").replace("\xa0", " ")
    texto = _ESPACOS_NA_PONTA.sub("\n", texto)
    texto = _ESPACOS_REPETIDOS.sub(" ", texto)
    return _LINHAS_EM_BRANCO.sub("\n\n", texto).strip()


# --- docx -----------------------------------------------------------------------------------------------------

def _docx(caminho: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(caminho))
    linhas: list[str] = []
    for filho in doc.element.body.iterchildren():  # na ordem do documento: parágrafos e tabelas misturados
        etiqueta = filho.tag.rsplit("}", 1)[-1]
        if etiqueta == "p":
            linhas.append(_paragrafo(Paragraph(filho, doc)))
        elif etiqueta == "tbl":
            linhas.append(_tabela(Table(filho, doc)))
    return "\n".join(linhas)


def _paragrafo(p) -> str:
    texto = p.text.strip()
    if not texto:
        return ""
    estilo = (p.style.name if p.style is not None else "") or ""
    nivel = re.match(r"(?:Heading|Título)\s*(\d)", estilo)
    if estilo == "Title" or estilo == "Título":
        return f"\n# {texto}\n"
    if nivel:
        return f"\n{'#' * min(int(nivel.group(1)), 3)} {texto}\n"
    if "List" in estilo or "Lista" in estilo:
        return f"- {texto}"
    return texto


def _tabela(tabela) -> str:
    linhas = []
    for linha in tabela.rows:
        celulas = []
        for celula in linha.cells:
            texto = " ".join(celula.text.split())
            if not celulas or celulas[-1] != texto:  # célula mesclada aparece repetida: uma vez só
                celulas.append(texto)
        if any(celulas):
            linhas.append("| " + " | ".join(celulas) + " |")
    return "\n" + "\n".join(linhas) + "\n" if linhas else ""


# --- pdf ------------------------------------------------------------------------------------------------------

def _pdf(caminho: Path) -> str:
    from pypdf import PdfReader

    leitor = PdfReader(str(caminho))
    if leitor.is_encrypted and not leitor.decrypt(""):
        raise ErroDeLeitura("PDF protegido por senha")
    paginas = []
    total = 0
    for pagina in leitor.pages[:MAX_PAGINAS_PDF]:  # uma página por vez: para no limite sem ler o resto do arquivo
        texto = (pagina.extract_text() or "").strip()
        if texto:
            paginas.append(texto)
            total += len(texto)
            if total > MAX_CARACTERES:
                break
    return "\n\n".join(paginas)


# --- planilhas ------------------------------------------------------------------------------------------------

def _xlsx(caminho: Path) -> str:
    from openpyxl import load_workbook

    with caminho.open("rb") as arquivo:  # o openpyxl recusa um caminho sem ".xlsx" no nome; o arquivo aberto, não
        return _abas(load_workbook(arquivo, read_only=True, data_only=True))


def _abas(pasta) -> str:
    try:
        partes = []
        for aba in pasta.worksheets:
            linhas = []
            for i, linha in enumerate(aba.iter_rows(values_only=True)):
                if i >= MAX_LINHAS_PLANILHA:
                    break
                celulas = ["" if c is None else str(c).strip() for c in linha]
                if any(celulas):
                    linhas.append(" | ".join(celulas).rstrip(" |"))
            if linhas:
                partes.append(f"## {aba.title}\n\n" + "\n".join(linhas))
        return "\n\n".join(partes)
    finally:
        pasta.close()


def _csv(caminho: Path) -> str:
    texto = caminho.read_bytes().decode("utf-8-sig", errors="replace")
    linhas = []
    for i, linha in enumerate(csv.reader(texto.splitlines())):
        if i >= MAX_LINHAS_PLANILHA:
            break
        celulas = [c.strip() for c in linha]
        if any(celulas):
            linhas.append(" | ".join(celulas).rstrip(" |"))
    return "\n".join(linhas)
